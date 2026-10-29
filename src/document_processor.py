"""
Document Processor module.
Handles extraction of text from PDF and DOCX files.
"""

import os
import logging
from pathlib import Path
from typing import Optional

from PyPDF2 import PdfReader
from docx import Document

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: str) -> str:
    """
    Extract all text content from a PDF file.

    Args:
        file_path: Path to the PDF file.

    Returns:
        Extracted text as a single string.

    Raises:
        FileNotFoundError: If the file does not exist.
        Exception: If the PDF cannot be read.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"PDF file not found: {file_path}")

    try:
        reader = PdfReader(file_path)
        text_parts = []
        for page_num, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
            else:
                logger.warning(f"No text extracted from page {page_num + 1} of {file_path}")
        full_text = "\n\n".join(text_parts)
        logger.info(f"Extracted {len(full_text)} characters from PDF: {file_path}")
        return full_text
    except Exception as e:
        logger.error(f"Error extracting text from PDF {file_path}: {e}")
        raise


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract all text content from a DOCX file.

    Args:
        file_path: Path to the DOCX file.

    Returns:
        Extracted text as a single string.

    Raises:
        FileNotFoundError: If the file does not exist.
        Exception: If the DOCX cannot be read.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    try:
        doc = Document(file_path)
        text_parts = []

        # Extract text from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)

        full_text = "\n\n".join(text_parts)
        logger.info(f"Extracted {len(full_text)} characters from DOCX: {file_path}")
        return full_text
    except Exception as e:
        logger.error(f"Error extracting text from DOCX {file_path}: {e}")
        raise


def extract_text(file_path: str) -> str:
    """
    Extract text from a document based on its file extension.

    Args:
        file_path: Path to the document file.

    Returns:
        Extracted text content.

    Raises:
        ValueError: If the file type is not supported.
    """
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Supported types: .pdf, .docx")


def get_document_metadata(file_path: str) -> dict:
    """
    Get metadata about a document file.

    Args:
        file_path: Path to the document file.

    Returns:
        Dictionary with file metadata.
    """
    path = Path(file_path)
    stat = path.stat()
    return {
        "filename": path.name,
        "extension": path.suffix.lower(),
        "size_bytes": stat.st_size,
        "size_readable": _format_file_size(stat.st_size),
    }


def _format_file_size(size_bytes: int) -> str:
    """Format file size in human-readable form."""
    for unit in ["B", "KB", "MB", "GB"]:
        if size_bytes < 1024:
            return f"{size_bytes:.1f} {unit}"
        size_bytes /= 1024
    return f"{size_bytes:.1f} TB"
